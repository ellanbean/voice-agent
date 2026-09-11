"""
Ops web app for the voice agent: upload leads (CSV/Excel), watch the board,
read call transcripts, work callbacks, manage do-not-call, pause campaigns.

    uvicorn web.app:app --reload --port 8080      # local
    Heroku: `web` process type (see heroku.yml)

Auth: Supabase Auth accounts (email + password). Roles live in `profiles`:
  admin           everything + user management
  team_lead       floor (live presence), reports, lead uploads, campaigns, DNC, board, calls, callbacks
  agent_sales     the console: receives handovers from the AI, sells, transfers to payments, wraps up
  agent_payments  the console: receives handovers from sales, takes payment, wraps up
  pending         no access (new signups) until an admin assigns a role
The first account to register becomes admin.

Env: WEB_SECRET (long random string), SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY
"""

from __future__ import annotations

import csv
import io
import os
import re
import secrets
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware

import db

load_dotenv()

HERE = Path(__file__).parent
app = FastAPI(title="into3 voice ops")
app.add_middleware(SessionMiddleware, secret_key=os.getenv("WEB_SECRET") or secrets.token_hex(32), https_only=False)
app.mount("/static", StaticFiles(directory=HERE / "static"), name="static")
templates = Jinja2Templates(directory=HERE / "templates")

E164 = re.compile(r"^\+[1-9]\d{6,14}$")
PENDING: dict[str, dict] = {}   # upload previews awaiting confirmation (server-side; cookies are too small)
LEAD_COLUMNS = ["phone", "name", "company", "country", "language", "source", "consent_at"]


# ------------------------------------------------------------------ auth & roles
# Roles: pending (no access) < agent_payments / agent_sales (console) < team_lead (floor, reports, leads, DNC) < admin (everything)
ROLE_RANK = {"pending": 0, "agent_payments": 1, "agent_sales": 1, "team_lead": 2, "admin": 3}
AGENT_ROLES = {"agent_sales", "agent_payments"}


def current_user(request: Request) -> dict:
    user = request.session.get("user")
    if not user:
        raise HTTPException(status_code=307, headers={"Location": "/login"})
    return user


def require(min_role: str):
    """Dependency: signed in AND role rank >= min_role's rank. 'agent' = any agent role or above."""
    need = 1 if min_role == "agent" else ROLE_RANK[min_role]
    def dep(request: Request) -> dict:
        user = current_user(request)
        if ROLE_RANK.get(user.get("role"), 0) < need:
            raise HTTPException(status_code=403, detail="Your account doesn't have access to this page yet. Ask an admin to set your role.")
        return user
    return dep


def _session_user(profile: dict) -> dict:
    return {"id": profile["id"], "email": profile["email"], "role": profile["role"],
            "name": profile.get("full_name") or "", "language": profile.get("language") or "en"}


@app.exception_handler(HTTPException)
async def http_error(request: Request, exc: HTTPException):
    if exc.status_code == 307 and exc.headers:
        return RedirectResponse(exc.headers["Location"], status_code=307)
    if exc.status_code in (403, 404) and "text/html" in request.headers.get("accept", "") and not request.url.path.startswith("/api/"):
        return templates.TemplateResponse(request, "error.html", {"code": exc.status_code, "detail": exc.detail}, status_code=exc.status_code)
    return JSONResponse({"detail": exc.detail}, status_code=exc.status_code, headers=exc.headers)


@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request):
    return templates.TemplateResponse(request, "login.html", {"error": None, "mode": "login"})


@app.post("/login")
def login(request: Request, email: str = Form(...), password: str = Form(...)):
    profile = db.sign_in(email.strip().lower(), password)
    if not profile:
        return templates.TemplateResponse(request, "login.html", {"error": "Wrong email or password", "mode": "login"}, status_code=401)
    request.session["user"] = _session_user(profile)
    return RedirectResponse("/console" if profile["role"] in AGENT_ROLES else "/", status_code=303)


@app.get("/signup", response_class=HTMLResponse)
def signup_form(request: Request):
    return templates.TemplateResponse(request, "login.html", {"error": None, "mode": "signup"})


@app.post("/signup")
def signup(request: Request, email: str = Form(...), password: str = Form(...), full_name: str = Form(""),
           phone: str = Form(""), employee_code: str = Form("")):
    if len(password) < 8:
        return templates.TemplateResponse(request, "login.html", {"error": "Password must be at least 8 characters", "mode": "signup"}, status_code=400)
    if len(full_name.strip()) < 2:
        return templates.TemplateResponse(request, "login.html", {"error": "Please enter your full name", "mode": "signup"}, status_code=400)
    try:
        db.sign_up(email.strip().lower(), password, full_name.strip(), phone.strip(), employee_code.strip())
    except Exception as e:  # noqa: BLE001 — duplicate email etc.
        return templates.TemplateResponse(request, "login.html", {"error": f"Could not create account: {e}", "mode": "signup"}, status_code=400)
    profile = db.sign_in(email.strip().lower(), password)
    request.session["user"] = _session_user(profile)
    return RedirectResponse("/", status_code=303)


@app.get("/logout")
def logout(request: Request):
    u = request.session.get("user")
    if u and u.get("role") in AGENT_ROLES:
        try:
            db.set_status(u["id"], "offline", None)
        except Exception:  # noqa: BLE001
            pass
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


# ------------------------------------------------------------------ users (admin)
@app.get("/users", response_class=HTMLResponse)
def users(request: Request, user: dict = Depends(require("team_lead"))):
    return templates.TemplateResponse(request, "users.html", {"users": db.profiles(), "roles": db.ROLES, "langs": db.LANGS, "me": user, "is_admin": user["role"] == "admin"})


@app.post("/users/{user_id}/priority")
def users_priority(user_id: str, priority: int = Form(...), user: dict = Depends(require("team_lead"))):
    db.set_priority(user_id, priority)
    return RedirectResponse("/users", status_code=303)


@app.post("/users/{user_id}/role")
def users_role(user_id: str, role: str = Form(...), language: str = Form("en"), user: dict = Depends(require("admin"))):
    if user_id == user["id"] and role != "admin":
        raise HTTPException(403, "You can't remove your own admin role.")
    db.set_role(user_id, role)
    db.set_language(user_id, language)
    return RedirectResponse("/users", status_code=303)


@app.post("/users/{user_id}/delete")
def users_delete(user_id: str, user: dict = Depends(require("admin"))):
    if user_id == user["id"]:
        raise HTTPException(403, "You can't delete yourself.")
    db.delete_user(user_id)
    return RedirectResponse("/users", status_code=303)


# ------------------------------------------------------------------ board
@app.get("/", response_class=HTMLResponse)
def board(request: Request, campaign: str = "", status: str = "", user: dict = Depends(require("agent"))):
    rows = db.lead_board(campaign=campaign or None, status=status or None)
    return templates.TemplateResponse(
        request, "board.html",
        {
            "rows": rows,
            "counts": db.status_counts(campaign or None),
            "campaigns": db.campaigns(),
            "campaign": campaign,
            "status": status,
        },
    )


@app.post("/campaigns/{name}/pause")
def pause_campaign(name: str, paused: str = Form(...), user: dict = Depends(require("team_lead"))):
    db.set_campaign_paused(name, paused == "1")
    return RedirectResponse("/", status_code=303)


# ------------------------------------------------------------------ upload
def _rows_from_upload(filename: str, data: bytes) -> list[dict]:
    """CSV or XLSX → list of dicts keyed by lower-cased header."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(it)]
        out = []
        for values in it:
            if not any(v not in (None, "") for v in values):
                continue
            out.append({header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(values) if i < len(header)})
        return out
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip().lower(): (v or "").strip() for k, v in row.items()} for row in reader]


def _normalise_phone(p: str) -> str:
    p = re.sub(r"[\s\-().]", "", p or "")
    if p.startswith("00"):
        p = "+" + p[2:]
    return p


def _validate(rows: list[dict], consent_all: bool) -> tuple[list[dict], list[dict]]:
    good, bad = [], []
    now = datetime.now(timezone.utc).isoformat()
    seen: set[str] = set()
    for i, r in enumerate(rows, start=2):
        phone = _normalise_phone(r.get("phone", ""))
        country = (r.get("country") or "").strip().upper()
        problems = []
        if not E164.match(phone):
            problems.append("phone not E.164 (+44…)")
        if len(country) != 2:
            problems.append("country must be ISO-2")
        if phone in seen:
            problems.append("duplicate in file")
        consent = (r.get("consent_at") or "").strip() or (now if consent_all else "")
        if not consent:
            problems.append("no consent_at (tick the box if the whole list is opted-in)")
        row = {
            "line": i, "phone": phone, "name": r.get("name") or "", "company": r.get("company") or "",
            "country": country, "language": (r.get("language") or "").lower(), "source": r.get("source") or "",
            "consent_at": consent, "problems": problems,
        }
        (bad if problems else good).append(row)
        seen.add(phone)
    return good, bad


@app.get("/upload", response_class=HTMLResponse)
def upload_form(request: Request, user: dict = Depends(require("team_lead"))):
    return templates.TemplateResponse(request, "upload.html", {"campaigns": db.campaigns(), "preview": None})


@app.post("/upload/preview", response_class=HTMLResponse)
async def upload_preview(
    request: Request, file: UploadFile = File(...), campaign: str = Form("general"),
    source: str = Form(""), consent_all: str = Form(""), user: dict = Depends(require("team_lead")),
):
    rows = _rows_from_upload(file.filename or "", await file.read())
    if source:
        for r in rows:
            r.setdefault("source", "")
            r["source"] = r["source"] or source
    good, bad = _validate(rows, consent_all == "1")
    token = secrets.token_urlsafe(16)
    PENDING[token] = {"campaign": campaign.strip() or "general", "good": good}
    request.session["pending"] = token
    return templates.TemplateResponse(
        request, "upload.html",
        {"campaigns": db.campaigns(), "preview": {"good": good, "bad": bad, "campaign": campaign, "filename": file.filename}},
    )


@app.post("/upload/commit")
def upload_commit(request: Request, user: dict = Depends(require("team_lead"))):
    token = request.session.pop("pending", None)
    pending = PENDING.pop(token, None) if token else None
    if not pending:
        return RedirectResponse("/upload", status_code=303)
    campaign = pending["campaign"]
    db.ensure_campaign(campaign)
    n = 0
    for r in pending["good"]:
        db.upsert_lead(
            {
                "phone": r["phone"], "name": r["name"] or None, "company": r["company"] or None,
                "country": r["country"], "language": r["language"] or None, "source": r["source"] or None,
                "consent_at": r["consent_at"], "campaign": campaign,
            }
        )
        n += 1
    return RedirectResponse(f"/?campaign={campaign}&imported={n}", status_code=303)


# ------------------------------------------------------------------ calls
@app.get("/calls", response_class=HTMLResponse)
def calls(request: Request, user: dict = Depends(require("agent"))):
    return templates.TemplateResponse(request, "calls.html", {"calls": db.recent_calls()})


@app.get("/calls/{call_id}", response_class=HTMLResponse)
def call_detail(request: Request, call_id: str, user: dict = Depends(require("agent"))):
    call = db.get_call(call_id)
    if not call:
        raise HTTPException(404)
    turns = []
    for item in (call.get("transcript") or {}).get("items", []):
        if item.get("type") == "message":
            content = item.get("content")
            text = " ".join(c if isinstance(c, str) else c.get("text", "") for c in content) if isinstance(content, list) else str(content or "")
            turns.append({"role": item.get("role"), "text": text})
        elif item.get("type") == "function_call":
            turns.append({"role": "tool", "text": f"{item.get('name')}({item.get('arguments')})"})
    segs = db.interp_segments(call_id) if call.get("interp_stats") or call.get("language") else []
    return templates.TemplateResponse(request, "call.html", {"call": call, "turns": turns, "interp": segs})


# ------------------------------------------------------------------ callbacks
@app.get("/callbacks", response_class=HTMLResponse)
def callbacks(request: Request, mine: str = "", user: dict = Depends(require("agent"))):
    items = db.my_callbacks(user["id"]) if mine else db.open_callbacks()
    return templates.TemplateResponse(request, "callbacks.html", {"items": items, "mine": bool(mine), "people": db.profiles()})


@app.post("/callbacks/{cb_id}/assign")
def callback_assign(cb_id: str, user_id: str = Form(""), user: dict = Depends(require("agent"))):
    db.assign_callback(cb_id, user_id or None)
    return RedirectResponse("/callbacks", status_code=303)


@app.post("/callbacks/{cb_id}/done")
def callback_done(cb_id: str, user: dict = Depends(require("agent"))):
    db.set_callback_done(cb_id, True)
    return RedirectResponse("/callbacks", status_code=303)


# ------------------------------------------------------------------ do-not-call
@app.get("/dnc", response_class=HTMLResponse)
def dnc(request: Request, user: dict = Depends(require("agent"))):
    return templates.TemplateResponse(request, "dnc.html", {"items": db.dnc_list(), "error": None})


@app.post("/dnc", response_class=HTMLResponse)
def dnc_add(request: Request, phone: str = Form(...), reason: str = Form("manual"), user: dict = Depends(require("team_lead"))):
    phone = _normalise_phone(phone)
    if not E164.match(phone):
        return templates.TemplateResponse(request, "dnc.html", {"items": db.dnc_list(), "error": f"{phone!r} is not E.164"})
    db.add_dnc(phone, reason, "manual")
    return RedirectResponse("/dnc", status_code=303)


@app.post("/dnc/remove")
def dnc_remove(phone: str = Form(...), user: dict = Depends(require("team_lead"))):
    db.remove_dnc(phone)
    return RedirectResponse("/dnc", status_code=303)


@app.get("/health")
def health():
    return {"ok": True}


# ================================================================== agent console
import asyncio as _asyncio
import json as _json
from livekit import api as lk_api


def _lk_token(identity: str, name: str, room: str, metadata: dict | None = None) -> str:
    tok = (
        lk_api.AccessToken()
        .with_identity(identity)
        .with_name(name)
        .with_grants(lk_api.VideoGrants(room_join=True, room=room, can_publish=True, can_subscribe=True, can_publish_data=True))
    )
    if metadata:
        tok = tok.with_metadata(_json.dumps(metadata))      # the interpreter reads {"language": "hi"} from here
    return tok.to_jwt()


# ---- interpreter: dispatched into the room when the customer's language differs from the agent's
INTERPRETER_ENABLED = os.getenv("INTERPRETER_ENABLED", "1") == "1"
INTERPRETER_AGENT_NAME = os.getenv("INTERPRETER_AGENT_NAME", "interpreter")
COUNTRY_LANG = {"GB": "en", "IE": "en", "DE": "de", "AT": "de", "CH": "de", "FR": "fr", "BE": "nl", "ES": "es", "IT": "it",
                "PT": "pt", "NL": "nl", "PL": "pl", "SE": "sv", "DK": "da", "NO": "no", "FI": "fi", "CZ": "cs", "RO": "ro",
                "HU": "hu", "GR": "el", "IN": "hi", "US": "en"}


def _customer_language(call: dict, lead: dict | None) -> str:
    """Daniel records the language he heard; fall back to the lead's language, then the country's."""
    lang = (call.get("language") or (lead or {}).get("language") or "").lower()[:2]
    if not lang and lead:
        lang = COUNTRY_LANG.get((lead.get("country") or "").upper(), "")
    return lang or "en"


async def _ensure_interpreter(room: str, customer_lang: str, agent_identity: str, agent_lang: str) -> bool:
    """Create the interpreter dispatch for this room once; later agents are picked up from their participant metadata."""
    lk = lk_api.LiveKitAPI()
    try:
        existing = await lk.agent_dispatch.list_dispatch(room_name=room)
        if any(d.agent_name == INTERPRETER_AGENT_NAME for d in existing):
            return False
        meta = {"customer_identity": "callee", "customer_language": customer_lang, "agents": {agent_identity: agent_lang}}
        await lk.agent_dispatch.create_dispatch(lk_api.CreateAgentDispatchRequest(agent_name=INTERPRETER_AGENT_NAME, room=room, metadata=_json.dumps(meta)))
        return True
    finally:
        await lk.aclose()


async def _send_control(room: str, payload: dict) -> None:
    lk = lk_api.LiveKitAPI()
    try:
        await lk.room.send_data(lk_api.SendDataRequest(room=room, data=_json.dumps(payload).encode(), kind=lk_api.DataPacket.RELIABLE,
                                                       destination_identities=[INTERPRETER_AGENT_NAME], topic="interp.control"))
    except Exception:  # noqa: BLE001 — interpreter may not be in the room
        pass
    finally:
        await lk.aclose()


async def _delete_room(room: str) -> None:
    lk = lk_api.LiveKitAPI()
    try:
        await lk.room.delete_room(lk_api.DeleteRoomRequest(room=room))
    except Exception:  # noqa: BLE001 — already gone
        pass
    finally:
        await lk.aclose()


async def _remove_participant(room: str, identity: str) -> None:
    lk = lk_api.LiveKitAPI()
    try:
        await lk.room.remove_participant(lk_api.RoomParticipantIdentity(room=room, identity=identity))
    except Exception:  # noqa: BLE001
        pass
    finally:
        await lk.aclose()


@app.get("/console", response_class=HTMLResponse)
def console(request: Request, user: dict = Depends(require("agent"))):
    me = db.get_profile(user["id"]) or {}
    return templates.TemplateResponse(request, "console.html", {"me": me, "livekit_url": os.getenv("LIVEKIT_URL", "")})


@app.post("/api/agent/status")
def api_agent_status(payload: dict, user: dict = Depends(require("agent"))):
    status = payload.get("status")
    if status not in ("available", "offline", "wrap_up"):
        raise HTTPException(400, "bad status")
    db.set_status(user["id"], status, None)
    return {"ok": True, "status": status}


@app.get("/api/agent/inbox")
def api_agent_inbox(user: dict = Depends(require("agent"))):
    """Polled by the console every 2s: my presence + any ringing transfer for me."""
    me = db.get_profile(user["id"]) or {}
    t = db.ringing_transfer_for(user["id"])
    return {"status": me.get("status"), "current_room": me.get("current_room"), "transfer": t}


@app.post("/api/transfers/{transfer_id}/accept")
async def api_transfer_accept(transfer_id: str, user: dict = Depends(require("agent"))):
    t = db.get_transfer(transfer_id)
    if not t or t["to_agent_id"] != user["id"] or t["status"] != "ringing":
        raise HTTPException(409, "transfer no longer ringing")
    me = db.get_profile(user["id"]) or {}
    agent_lang = (me.get("language") or "en")[:2]
    ctx = db.call_context(t["room"])
    customer_lang = _customer_language(ctx.get("call") or {}, ctx.get("lead"))
    identity = f"human-{user['id'][:8]}"           # "agent-…" is what LiveKit gives AI workers; keep humans distinct

    db.set_transfer_status(transfer_id, "accepted")
    db.set_status(user["id"], "on_call", t["room"])
    db.set_call_stage(t["room"], t["to_stage"], language=customer_lang)
    db.open_segment(t["room"], t["to_stage"], user["id"])

    interp = {"enabled": False, "identity": INTERPRETER_AGENT_NAME, "customer_language": customer_lang, "my_language": agent_lang}
    if INTERPRETER_ENABLED and customer_lang != agent_lang:
        interp["enabled"] = True
        try:
            await _ensure_interpreter(t["room"], customer_lang, identity, agent_lang)
        except Exception as e:  # noqa: BLE001 — still connect the agent; the console shows the warning
            interp["enabled"] = False
            interp["error"] = f"interpreter unavailable: {e}"
    token = _lk_token(identity, user.get("name") or user["email"], t["room"], {"language": agent_lang, "role": t["to_stage"]})
    return {"ok": True, "room": t["room"], "token": token, "url": os.getenv("LIVEKIT_URL", ""), "context": ctx, "interpreter": interp}


# ---- simulator: a browser plays the customer so the whole flow runs without a SIP trunk
@app.get("/sim", response_class=HTMLResponse)
def sim_page(request: Request, user: dict = Depends(require("team_lead"))):
    return templates.TemplateResponse(request, "sim.html", {"livekit_url": os.getenv("LIVEKIT_URL", "")})


@app.post("/api/sim/start")
async def api_sim_start(payload: dict, user: dict = Depends(require("team_lead"))):
    """Create a test lead, dispatch Daniel into a fresh room in sim mode, and hand the browser a 'callee' token."""
    name = (payload.get("name") or "Test Customer").strip()
    country = (payload.get("country") or "DE").upper()[:2]
    language = (payload.get("language") or COUNTRY_LANG.get(country, "en")).lower()[:2]
    phone = f"+000{secrets.randbelow(10**9):09d}"                     # unique fake number per run
    lead = db.upsert_lead({"phone": phone, "name": name, "country": country, "language": language, "source": "simulator",
                           "consent_at": datetime.now(timezone.utc).isoformat(), "campaign": "sim", "status": "calling"})
    room = f"sim-{secrets.token_hex(4)}"
    meta = {"sim": True, "phone": phone, "lead_id": lead["id"], "name": name, "country": country, "language": language,
            "source": "simulator", "campaign": "sim", "company": payload.get("company") or ""}
    lk = lk_api.LiveKitAPI()
    try:
        await lk.agent_dispatch.create_dispatch(lk_api.CreateAgentDispatchRequest(agent_name=os.getenv("AGENT_NAME", "sales-caller"), room=room, metadata=_json.dumps(meta)))
    finally:
        await lk.aclose()
    token = _lk_token("callee", name, room)
    return {"ok": True, "room": room, "token": token, "url": os.getenv("LIVEKIT_URL", ""), "lead": lead}


@app.post("/api/sim/{room}/hangup")
async def api_sim_hangup(room: str, user: dict = Depends(require("team_lead"))):
    await _delete_room(room)
    try:
        db.end_call(room)
    except Exception:  # noqa: BLE001
        pass
    return {"ok": True}


@app.post("/api/calls/{room}/customer-language")
async def api_set_customer_language(room: str, payload: dict, user: dict = Depends(require("agent"))):
    """Agent corrects the detected language mid-call (e.g. Daniel heard Dutch, it's actually German)."""
    lang = (payload.get("language") or "").lower()[:2]
    if len(lang) != 2:
        raise HTTPException(400, "language must be an ISO code like de")
    db.set_call_language(room, lang)
    await _send_control(room, {"op": "set_customer_language", "language": lang})
    return {"ok": True, "language": lang}


@app.get("/api/calls/{room}/interp-stats")
def api_interp_stats(room: str, user: dict = Depends(require("agent"))):
    call = db.get_call_by_room(room) or {}
    return {"stats": call.get("interp_stats"), "segments": db.interp_segments(call["id"]) if call.get("id") else []}


@app.post("/api/transfers/{transfer_id}/decline")
def api_transfer_decline(transfer_id: str, user: dict = Depends(require("agent"))):
    t = db.get_transfer(transfer_id)
    if t and t["to_agent_id"] == user["id"] and t["status"] == "ringing":
        db.set_transfer_status(transfer_id, "declined")
    db.set_status(user["id"], "available", None)
    return {"ok": True}


@app.get("/api/calls/by-room/{room}")
def api_call_context(room: str, user: dict = Depends(require("agent"))):
    return db.call_context(room)


@app.post("/api/calls/{room}/transfer-to-payments")
async def api_transfer_payments(room: str, payload: dict, user: dict = Depends(require("agent"))):
    """Sales agent hands the customer to a payments agent. Rings one; the sales
    agent stays on the line until the payments agent has joined, then leaves."""
    summary = (payload.get("summary") or "").strip()
    agent = db.pick_available_agent("agent_payments")
    if not agent:
        raise HTTPException(409, "No payments agent is available right now.")
    t = db.create_transfer(room, "sales", "payments", agent["id"], summary)
    db.set_status(agent["id"], "ringing", room)
    return {"ok": True, "transfer_id": t["id"], "agent": agent.get("full_name") or agent.get("email")}


@app.get("/api/transfers/{transfer_id}")
def api_transfer_get(transfer_id: str, user: dict = Depends(require("agent"))):
    t = db.get_transfer(transfer_id)
    if not t:
        raise HTTPException(404)
    return {"status": t["status"]}


@app.post("/api/transfers/{transfer_id}/cancel")
def api_transfer_cancel(transfer_id: str, user: dict = Depends(require("agent"))):
    t = db.get_transfer(transfer_id)
    if t and t["status"] == "ringing":
        db.set_transfer_status(transfer_id, "cancelled")
        if t.get("to_agent_id"):
            db.set_status(t["to_agent_id"], "available", None)
    return {"ok": True}


@app.post("/api/calls/{room}/leave")
async def api_leave(room: str, user: dict = Depends(require("agent"))):
    """Agent leaves the room without ending the call (after handing to payments)."""
    db.close_segment(room, "sales", user["id"])
    db.set_status(user["id"], "wrap_up", room)
    return {"ok": True}


@app.post("/api/calls/{room}/hangup")
async def api_hangup(room: str, user: dict = Depends(require("agent"))):
    """Agent ends the call for everyone."""
    await _delete_room(room)
    db.end_call(room)
    db.set_status(user["id"], "wrap_up", room)
    return {"ok": True}


@app.post("/api/calls/{room}/wrapup")
def api_wrapup(room: str, payload: dict, user: dict = Depends(require("agent"))):
    """Mandatory after every human segment: disposition + remarks (+ sale amount)."""
    disposition = payload.get("disposition")
    if disposition not in ("sale", "callback", "not_interested", "no_decision", "wrong_person", "dropped", "other"):
        raise HTTPException(400, "bad disposition")
    stage = "payments" if user.get("role") == "agent_payments" else "sales"
    wrap = {"disposition": disposition, "remarks": (payload.get("remarks") or "").strip()}
    if disposition == "sale":
        wrap["sale_amount"] = float(payload.get("sale_amount") or 0)
        wrap["currency"] = (payload.get("currency") or "EUR").upper()
    # close this agent's segment if still open, then record
    call = db.get_call_by_room(room)
    if call:
        db.client().table("call_segments").update({**wrap, "ended_at": db._now()}).eq("call_id", call["id"]).eq("agent_id", user["id"]).is_("ended_at", "null").execute()
        db.client().table("call_segments").update(wrap).eq("call_id", call["id"]).eq("agent_id", user["id"]).is_("disposition", "null").execute()
        if disposition == "sale":
            db.client().table("calls").update({"status": "converted"}).eq("room", room).execute()
            if call.get("lead_id"):
                db.client().table("leads").update({"status": "converted", "last_outcome": "sale"}).eq("id", call["lead_id"]).execute()
        elif disposition == "not_interested" and call.get("phone"):
            db.add_dnc(call["phone"], "human agent: not interested", "agent")
            if call.get("lead_id"):
                db.client().table("leads").update({"status": "not_interested", "last_outcome": "not_interested"}).eq("id", call["lead_id"]).execute()
        elif disposition == "callback":
            db.add_callback(call.get("lead_id"), room, payload.get("callback_when") or "as agreed", wrap["remarks"])
            if call.get("lead_id"):
                db.client().table("leads").update({"status": "callback", "last_outcome": "callback"}).eq("id", call["lead_id"]).execute()
    db.set_status(user["id"], "available", None)
    return {"ok": True}


@app.post("/webhooks/livekit")
async def livekit_webhook(request: Request):
    """LiveKit Cloud → Settings → Webhooks → https://<app>/webhooks/livekit.
    room_finished closes the call record when the customer hangs up on a human."""
    body = await request.body()
    auth = request.headers.get("Authorization", "")
    try:
        receiver = lk_api.WebhookReceiver(lk_api.TokenVerifier())
        event = receiver.receive(body.decode(), auth)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(401, f"bad webhook signature: {e}")
    if event.event == "room_finished" and event.room and event.room.name.startswith("call-"):
        db.end_call(event.room.name)
    elif event.event == "participant_left" and event.participant and event.participant.identity == "callee" and event.room:
        # customer hung up while a human was on: end the call for everyone
        await _delete_room(event.room.name)
        db.end_call(event.room.name, final_status="callee_hung_up")
    return {"ok": True}


# ================================================================== team lead: floor + reports
@app.get("/floor", response_class=HTMLResponse)
def floor(request: Request, user: dict = Depends(require("team_lead"))):
    return templates.TemplateResponse(request, "floor.html", {"rows": db.floor()})


@app.get("/api/floor")
def api_floor(user: dict = Depends(require("team_lead"))):
    return db.floor()


@app.get("/reports", response_class=HTMLResponse)
def reports(request: Request, days: int = 7, user: dict = Depends(require("team_lead"))):
    return templates.TemplateResponse(request, "reports.html", {"rows": db.agent_report(days), "days": days, "interp": db.interp_report(days)})
